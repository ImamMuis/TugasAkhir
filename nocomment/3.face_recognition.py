import os
import cv2
import time
import busio
import socket
import telepot
import datetime
import openpyxl 
import RPi.GPIO as GPIO
from board import SCL, SDA
from adafruit_pca9685 import PCA9685
from adafruit_servokit import ServoKit

pin_solenoid    = 14
pin_pintuBuka   = 22
pin_pintuTutup  = 27
pin_sensorPIR   = 17
pin_motorLogic1 = 15
pin_motorLogic2 = 18

cam = cv2.VideoCapture(0)
resolution = 480

fps = 25
recordTime = 10.4
faceCompare = 20
motorMIN_Volt = 5
motorMAX_Volt = 24
servoX_Degree = 90
servoY_Degree = 90
motorPWM_Channel = 3
waktuPintuTerbuka = 5
faceDetect_Interval = 3
telebotAdminID  = 1338050139
facePosition = [250, 390, 170, 310]

root = '/home/pi/1.TugasAkhir/'
userDir = root + 'img_record'
teleBot_PWD = 'Mkt2k21'
fileUser = root + 'data/Username.txt'
dataMasuk = root + 'data/userMasuk.xlsx'
tokenBot = '2077135119:AAF3srU5w3127y0CNkDUJeMGJu2-3aTVQYI'

ratio = 4 / 3
rgbHeight = resolution
rgbWidth = int(round(rgbHeight * ratio, 0))
cam.set(3, rgbWidth)
cam.set(4, rgbHeight)

i2c_bus = busio.I2C(SCL, SDA)
pca = PCA9685(i2c_bus)
pca.frequency = 50

kit = ServoKit(channels=16)
kit.servo[0].set_pulse_width_range(500, 2750)
kit.servo[1].set_pulse_width_range(500, 2750)

with open(fileUser) as user:
    names = user.read().splitlines()

font = cv2.FONT_HERSHEY_SIMPLEX
cascadePath = root + 'data/haarcascade_frontalface_default.xml'
faceDetector = cv2.CascadeClassifier(cascadePath)
faceRecognizer = cv2.face.LBPHFaceRecognizer_create()
faceRecognizer.read(root + 'data/trainer.xml')

DATE = ''
Time = ''
command = ''
videoOut = ''

Quit = False
openDoor = False
QuitFlag = False
faceState = False
recordFlag = False

imgRGB = 0
chat_id = 0
timeNow = 0
stepDoor = 0
videoStep = 0
motorZERO = 0
frameRecord = 0
countFaceCompare = 0
DetectedFace_Last =  0
notDetectedTime_Now = 0
notDetectedTime_Last = 0
totalUser = len(names)
countID = [0] * totalUser
frameTotal = int(recordTime * fps)
detectResult = [0] * faceCompare

day = {
    'Sun' : 'Minggu',
    'Mon' : 'Senin',
    'Tue' : 'Selasa',
    'Wed' : 'Rabu',
    'Thu' : 'Kamis',
    'Fri' : 'Jumat',
    'Sat' : 'Sabtu',
}

month = {
    'Jan' : 'Januari',
    'Feb' : 'Februari',
    'Mar' : 'Maret',
    'Apr' : 'April',
    'Mei' : 'Mei',
    'Jun' : 'Juni',
    'Jul' : 'Juli',
    'Agt' : 'Agustus',
    'Sep' : 'September',
    'Oct' : 'Oktober',
    'Nov' : 'Nopember',
    'Dec' : 'Desember'
}

GPIO.setmode(GPIO.BCM)
GPIO.setup(pin_sensorPIR, GPIO.IN)
GPIO.setup(pin_pintuBuka, GPIO.IN)
GPIO.setup(pin_pintuTutup, GPIO.IN)
GPIO.setup(pin_solenoid, GPIO.OUT)
GPIO.setup(pin_motorLogic1, GPIO.OUT)
GPIO.setup(pin_motorLogic2, GPIO.OUT)
GPIO.output(pin_solenoid, 0)
GPIO.output(pin_motorLogic1, 0)
GPIO.output(pin_motorLogic2, 0)

time.sleep(0.1)

def Selisih(current, prev):
    num = float(current) - float(prev)
    if num < 0:
        num += 60
    num = round(num, 3)
    return num

def getCurrent(data):
    global day
    global month

    now = datetime.datetime.now()
    if data == 'DATE':
        dayNow   = day[now.strftime('%a')]
        monthNow = month[now.strftime('%b')]
        date  = now.strftime('%d')
        year  = now.strftime('%Y')
        value = f'{dayNow}, {date} {monthNow} {year}'

    elif data == 'Date':
        value = now.strftime('%Y-%m-%d.%H.%M.%S')

    elif data == 'Time':
        Time  = now.strftime('%H:%M:%S.')
        sec   = str(round(float(str(now).split(':')[-1]), 3))
        value = Time + sec.split('.')[-1]

    elif data == 'second':
        value = str(round(float(str(now).split(':')[-1]), 3))

    else:
        print('Parameter Salah!')

    return value

def saveToExcel(user, dates, time):
    wb = openpyxl.load_workbook(dataMasuk) 
    sheet = wb.active 
    maxRow = sheet.max_row
    day = dates.split(',')[0]
    date = dates.split(',')[1]
    newdata = [[maxRow, user, day, date, time]]
    
    for row in newdata:
        sheet.append(row)
    
    wb.save(dataMasuk)
    return maxRow

def saveImage(img, date, time):
    waktu = getCurrent('Date')
    img = cv2.putText(img, time, (15, 447), font, 0.5, (54, 67, 244), 2)
    img = cv2.putText(img, date, (15, 462), font, 0.5, (54, 67, 244), 2)
    namaFile = 'Snapshot.' + waktu + '.jpg'
    cv2.imwrite(userDir + '/' + namaFile, img)

def sendImage(chatId, msgCaption=None):
    file = os.listdir(userDir)
    lastImage = sorted(file, key = lambda x: os.path.splitext(x)[0])
    foto = lastImage[-1]
    file = userDir + '/' + foto
    fileType = file.split('.')[-1]

    if fileType == 'jpg':
        bot.sendPhoto(chatId, photo=open(file, 'rb'),
                      caption=msgCaption)

    elif fileType == 'mp4':
        bot.sendVideo(chatId, video=open(file, 'rb'), 
                      caption=msgCaption)

def servoMove(channel, degree):
    if degree < 0:
        print('Range terlalu kecil!')
        kit.servo[channel].angle = 0

    elif degree > 180:
        print('Range terlalu besar!')
        kit.servo[channel].angle = 180
    
    else:
        print(f'Servo {channel} : {degree}°')
        kit.servo[channel].angle = degree

def recordVideo(img):
    global DATE
    global Time
    global chat_id
    global command
    global timeNow
    global videoOut
    global videoStep
    global recordFlag
    global frameRecord
    global servoX_Degree
    global servoY_Degree

    if videoStep == 0:
        servoX_Degree = 80
        servoY_Degree = 90
        servoMove(0, servoX_Degree)
        servoMove(1, servoY_Degree)
        videoStep += 1
        print('Create video')
        timeNow = time.time()
        Date = getCurrent('Date')
        videoName = f'{userDir}/Snapshot.{Date}.mp4'
        videoEncode = cv2.VideoWriter_fourcc(*'mp4v')
        videoOut = cv2.VideoWriter(videoName, videoEncode, fps, (640, 480))

    elif videoStep == 1:
        frameRecord += 1
        DATE = getCurrent('DATE')
        Time = getCurrent('Time')
        img = cv2.flip(img, 1)
        img = cv2.putText(img, Time, (15, 447), font, 0.5, (54, 67, 244), 2)
        img = cv2.putText(img, DATE, (15, 462), font, 0.5, (54, 67, 244), 2)
        videoOut.write(img)

        if frameRecord <= 60:
            servoY_Degree -= 1
            servoMove(1, servoY_Degree)

        elif frameRecord > 60 and frameRecord <= 190:
            servoY_Degree += 1
            servoMove(1, servoY_Degree)

        elif frameRecord > 190 and frameRecord <= 260:
            servoY_Degree -= 1
            servoMove(1, servoY_Degree)

        if frameRecord == frameTotal:
            videoStep += 1
            print('Record done')
            bot.sendMessage(chat_id, 
            'Perekaman selesai, video sedang dikirim...')

    elif videoStep == 2:
        servoX_Degree = 90
        servoY_Degree = 90
        servoMove(0, servoX_Degree)
        servoMove(1, servoY_Degree)
        videoOut.release()
        print('Video saved')
        telebotText = f'Monitoring Lingkungan\n{Time}\n{DATE}'
        sendImage(chat_id, telebotText)
        timeNow = 0
        command = ''
        videoStep = 0
        frameRecord = 0
        recordFlag = False

def setSpeed(speedMin, speedMax):
    Vmax = 24
    PWM_width = 2 ** 16 - 1
    dutyCycle = speedMin / Vmax
    motorMIN_PWM = int(round(dutyCycle * PWM_width, 0))
    dutyCycle = speedMax / Vmax
    motorMAX_PWM = int(round(dutyCycle * PWM_width, 0))

    return motorMIN_PWM, motorMAX_PWM

def motorSpeed(begin, end, step, accel):    
    if accel == 0:
        pca.channels[motorPWM_Channel].duty_cycle = motorMIN

    elif accel == 1:
        for i in range(begin, end, step):
            pca.channels[motorPWM_Channel].duty_cycle = i
            if GPIO.input(pin_pintuBuka) or GPIO.input(pin_pintuTutup) == 1:
                break

def motorStart(mode):
    if mode == 'FORWARD':
        print('Mode: Forward')
        GPIO.output(pin_motorLogic1, 1)
        GPIO.output(pin_motorLogic2, 0)
        motorSpeed(motorMIN, motorMAX, 200, 1)

    elif mode == 'REVERSE':
        print('Mode: Reverse')
        GPIO.output(pin_motorLogic1, 0)
        GPIO.output(pin_motorLogic2, 1)
        motorSpeed(motorMIN, motorMAX, 200, 1)

    elif mode == 'CLOSE':
        GPIO.output(pin_motorLogic1, 0)
        GPIO.output(pin_motorLogic2, 1)
        motorSpeed(motorMIN, 0, 0, 0)

    else:
        print('Parameter "mode" harus FORWARD atau REVERSE!')

def motorStop(forceBreak = 0):
    if forceBreak == 0:
        motorSpeed(motorMIN, motorZERO, -200, 1)

    elif forceBreak == 1:
        GPIO.output(pin_motorLogic1, 0)
        GPIO.output(pin_motorLogic2, 0)

    print('Motor Stop')

def setupPIR():
    print('Menyiapkan Sensor PIR...') 
    time.sleep(0.02)
    
    while GPIO.input(pin_sensorPIR) == 1:
        print('Sensor PIR belum siap')
        print('Mohon untuk tidak ada pergerakan terlebih dahulu!')
        time.sleep(0.5)

    print('Sensor PIR Siap!')
    time.sleep(0.02)

def setupPintu():
    print('Memastikan Pintu Tertutup')
    if GPIO.input(pin_pintuTutup) == 0:
        print('Pintu sedang terbuka! Pintu akan ditutup...')
        GPIO.output(pin_solenoid, 1)
        time.sleep(0.5)

        while GPIO.input(pin_pintuTutup) == 0:
            motorStart('CLOSE') 
                  
    motorStop(1)
    print('Pintu sudah tertutup!')
    GPIO.output(pin_solenoid, 0)

def sistemPintu(kondisi):
    if kondisi == 'Buka':
        print('Pintu dibuka')
        GPIO.output(pin_solenoid, 1)
        while GPIO.input(pin_pintuBuka) == 0:
            motorStart('FORWARD')
        
        print('Pintu terbuka')

    elif kondisi == 'Tutup':
        print('Pintu ditutup')
        while GPIO.input(pin_pintuTutup) == 0:
            motorStart('REVERSE')
 
        GPIO.output(pin_solenoid, 0)
        print('Pintu tertutup')
        
    motorStop(1) 

def teleBot(msg):
    global Quit
    global imgRGB
    global chat_id
    global command
    global QuitFlag
    global recordFlag
    global teleBot_PWD
    global servoX_Degree
    global servoY_Degree

    chat_id = msg['chat']['id']
    command = msg['text']
    
    print('From User : ', chat_id)
    print('Command   : ', command)

    disableCommnad = ['^', 'v', '<', '>', 'Reset Servo', 
                      'Stop Sistem', 'Rekam Video']

    show_keyboard = {'keyboard':[   
        [    'Ambil Foto'    ,      '^'     , 'Rekam Video'], 
        [        '<'         , 'Reset Servo',      '>'     ],
        ['History User Masuk',      'v'     , 'Stop Sistem']
                                ]}

    if command == '/start':
        bot.sendMessage(chat_id, 'Silakan pilih perintah:', 
                        reply_markup=show_keyboard)
        
    elif recordFlag == True and command in disableCommnad:
        bot.sendMessage(chat_id, 'Video sedang direkam!')

    elif command == 'Stop Sistem':
        QuitFlag = True
        bot.sendMessage(chat_id, 
                        'Masukan PIN untuk stop Sistem Face Recognition')

    elif command == teleBot_PWD:
        Quit = True
        bot.sendMessage(chat_id, 'Sistem Face Recognition terhenti...')

    elif command != teleBot_PWD and QuitFlag == True:
        QuitFlag = False
        bot.sendMessage(chat_id, 'Password Salah!')

    elif command == 'Ambil Foto':
        Date = getCurrent('DATE')
        Time = getCurrent('Time')
        saveImage(imgRGB, Date, Time)
        telebotText = f'Snapshot\n{Time}\n{Date}'
        sendImage(chat_id, telebotText)

    elif command == 'Rekam Video':
        recordFlag = True
        bot.sendMessage(chat_id, 'Monitoring Lingkungan akan dilakukan')

    elif command == 'History User Masuk':
        bot.sendDocument(chat_id, document=open(dataMasuk, 'rb'))

    elif command == '^':
        servoY_Degree -= 5
        servoMove(0, servoY_Degree)
        telebotText = f'Servo Y: {servoY_Degree}°'
        bot.sendMessage(chat_id, telebotText)

    elif command == 'v':
        servoY_Degree += 5
        servoMove(0, servoY_Degree)
        telebotText = f'Servo Y: {servoY_Degree}°'
        bot.sendMessage(chat_id, telebotText)

    elif command == '<':
        servoX_Degree -= 5
        servoMove(1, servoX_Degree)
        telebotText = f'Servo X: {servoX_Degree}°'
        bot.sendMessage(chat_id, telebotText)

    elif command == '>':
        servoX_Degree += 5
        servoMove(1, servoX_Degree)
        telebotText = f'Servo X: {servoX_Degree}°'
        bot.sendMessage(chat_id, telebotText)
    
    elif command == 'Reset Servo':
        servoX_Degree = 90
        servoY_Degree = 90
        servoMove(1, servoX_Degree)
        servoMove(0, servoY_Degree)
        value1 = f'Servo X: {servoX_Degree}°\n'
        value2 = f'Servo Y: {servoY_Degree}°'
        telebotText = value1 + value2
        bot.sendMessage(chat_id, telebotText)

    elif command == 'My id':
        bot.sendMessage(chat_id, str(chat_id))

    else:
        bot.sendMessage(chat_id, 'Input belum tersedia!')
        bot.sendMessage(chat_id, 'Silakan pilih perintah:', 
                        reply_markup=show_keyboard) 

def detectFace():
    global imgRGB
    global chat_id
    global timeNow
    global openDoor
    global stepDoor
    global faceState
    global recordFlag
    global servoY_Degree
    global servoX_Degree
    global countFaceCompare
    global waktuPintuTerbuka
    global DetectedFace_Last
    global notDetectedTime_Now
    global faceDetect_Interval

    jumlahWajah = 0
    jarakWajah = [75, 85]

    grayHeight = 165
    grayWidth = int(round(grayHeight * ratio, 0))
    scaling = rgbWidth / grayWidth
    succes, frame = cam.read()

    if succes:
        imgRGB = cv2.flip(frame, 1)
        imgGray = cv2.cvtColor(imgRGB, cv2.COLOR_BGR2GRAY)
        imgGray = cv2.resize(imgGray, (grayWidth, grayHeight))
        faces = faceDetector.detectMultiScale(imgGray, 1.2, 3)

        for x1, y1, w1, h1 in faces:
            x2 = int(round(x1 * scaling, 0))
            y2 = int(round(y1 * scaling, 0))
            w2 = int(round(w1 * scaling, 0))
            h2 = int(round(h1 * scaling, 0))
            cX = int(round(x2 + w2 / 2 , 0))
            cY = int(round(y2 + h2 / 2 , 0))

            jumlahWajah = int(str(faces.shape[0]))
            Id, confidence = faceRecognizer.predict(imgGray[y1:y1+h1, 
                                                    x1:x1+w1])
            imgRGB = cv2.rectangle(imgRGB, (x2, y2), (x2+w2, y2+h2), 
                                   (186, 39, 59), 2)

            if jumlahWajah > 1:
                imgRGB = cv2.putText(imgRGB, 'Wajah lebih dari satu!', (15, 50), 
                                     font, 0.7, (54, 67, 244), 2)
            
            elif jumlahWajah == 1:
                if w1 <= jarakWajah[0] :
                    imgRGB = cv2.putText(imgRGB, 'Wajah terlalu jauh!', (15, 25), 
                                         font, 0.7, (54, 67, 244), 2)
                
                elif w1 >= jarakWajah[1]:
                    imgRGB = cv2.putText(imgRGB, 'Wajah terlalu dekat!', (15, 25), 
                                         font, 0.7, (54, 67, 244), 2)

                if recordFlag == False:
                    if cX < facePosition[0]:
                        servoX_Degree += 1
                        servoMove(1, servoX_Degree)
                        imgRGB = cv2.putText(imgRGB, 'Wajah terlalu kiri!', 
                                     (15, 50), font, 0.7, (54, 67, 244), 2)
                
                    elif cX > facePosition[1]:
                        servoX_Degree -= 1
                        servoMove(1, servoX_Degree)
                        imgRGB = cv2.putText(imgRGB, 'Wajah terlalu kanan!', 
                                     (15, 75), font, 0.7, (54, 67, 244), 2)
                        
                    if cY < facePosition[2]:
                        servoY_Degree -= 1
                        servoMove(0, servoY_Degree)
                        imgRGB = cv2.putText(imgRGB, 'Wajah terlalu atas!', 
                                     (15, 50), font, 0.7, (54, 67, 244), 2)
                        
                    elif cY > facePosition[3]:
                        servoY_Degree += 1
                        servoMove(0, servoY_Degree)
                        imgRGB = cv2.putText(imgRGB, 'Wajah terlalu bawah!', 
                                     (15, 75), font, 0.7, (54, 67, 244), 2)
                
                if confidence >= 80 and confidence <= 100:
                    nameID = names[Id]
                    confidenceText = ' {0}%'.format(round(confidence))
                    
                else:
                    nameID = names[0]
                    confidenceText = ' {0}%'.format(round(100-confidence))

                imgRGB = cv2.putText(imgRGB, nameID, (x2, y2-5), 
                                     font, 0.7, (145, 202, 19), 2)
                imgRGB = cv2.putText(imgRGB, confidenceText, (x2+w2-60, y2+h2-5), 
                                     font, 0.7, (145, 202, 19), 2)
                
                if recordFlag == False:
                    if countFaceCompare < faceCompare:
                        detectResult[countFaceCompare] = names.index(nameID)
                        countFaceCompare += 1

        if recordFlag == True:
            recordVideo(imgRGB)

        cv2.imshow('Face Recognition', imgRGB)
        cv2.moveWindow('Face Recognition', 0, 0)
        DetectedFace_Now = jumlahWajah

        if countFaceCompare == faceCompare:
            for i in range(totalUser):
                countID[i] = detectResult.count(i)

            nameResult = names[countID.index(max(countID))]

            if DetectedFace_Last == 0 and DetectedFace_Now == 1:
                DetectedFace_Last = 1
                DetectedTime_Now = getCurrent('second')
                TimeBetween = Selisih(DetectedTime_Now, notDetectedTime_Now)
                
                if chat_id == 0:
                    chat_id = telebotAdminID

                if TimeBetween > faceDetect_Interval:
                    faceState = True
                    Date = getCurrent('DATE')
                    Time = getCurrent('Time')
                    count1 = saveToExcel(nameResult, Date, Time)

                    print('Deteksi ke     :', count1)
                    print('Pukul          :', Time)
                    print('Hari, Tanggal  :', Date)
                    print('Compare Wajah  :', detectResult)
                    print('User terdeteksi:', nameResult)
                    saveImage(imgRGB, Date, Time)

                    if nameResult in names[1:]:
                        stepDoor = 1
                        openDoor = True
                        telebotText  = f'User {nameResult} masuk\n'

                    else:
                        telebotText = 'Wajah tidak dikenali!\n'

                    Time = f'Pukul : {Time}\n'
                    Date = f'Tanggal : {Date}\n'
                    telebotText = telebotText + Time + Date
                    sendImage(chat_id, telebotText)

            elif DetectedFace_Last == 1 and DetectedFace_Now == 0:
                DetectedFace_Last   = 0
                notDetectedTime_Now = getCurrent('second')

            elif DetectedFace_Last == 0 and DetectedFace_Now == 0:
                notDetectedTime_Last = getCurrent('second')
                TimeBetween = Selisih(notDetectedTime_Last, notDetectedTime_Now)

                if TimeBetween > faceDetect_Interval and faceState == True:
                    faceState = False
                    countFaceCompare = 0
                    print('Tidak Ada Wajah!')
                    telebotText = 'Tidak ada wajah dalam ' + 
                                   faceDetect_Interval + ' detik terakhir'
                    bot.sendMessage(chat_id, telebotText)

        if openDoor == True:
            if stepDoor == 1:
                stepDoor += 1
                sistemPintu('Buka')
                timeNow = time.time()

            elif stepDoor == 2:
                selisih = time.time() - timeNow

                if selisih < waktuPintuTerbuka:
                    if GPIO.input(pin_sensorPIR) == 1:`
                        print('Anda sudah masuk') 
                        sistemPintu('Tutup')
                    
                elif selisih >= waktuPintuTerbuka:
                    stepDoor = 0
                    openDoor = False
                    print('Anda tidak segera masuk')
                    sistemPintu('Tutup')

try:
    socket.setdefaulttimeout(3)
    socket.socket(socket.AF_INET,
        socket.SOCK_STREAM).connect(('8.8.8.8', 53))
    print('Terkoneksi dengan internet')

    bot = telepot.Bot(tokenBot)
    bot.message_loop(teleBot)
    telebotText = 'Sistem Face Recognition Berjalan\n'
    print(telebotText)
    bot.sendMessage(telebotAdminID, telebotText)
    motorMIN, motorMAX = setSpeed(motorMIN_Volt, motorMAX_Volt)

    setupPIR()
    setupPintu()
    servoMove(0, 90)
    servoMove(1, 90)

    while True:
        detectFace()

        if cv2.waitKey(1) & 0xFF == ord('q') or Quit == True:
            break

except socket.error:
    print('Tidak terkoneksi dengan internet!')

except KeyboardInterrupt:
    print('Program Stop')
    
finally:
    cam.release()
    cv2.destroyAllWindows()
    setupPintu()
    servoMove(0, 90)
    servoMove(1, 90)
    GPIO.cleanup()